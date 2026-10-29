# c:\Users\han\development\antigraviy\unfi-ing\generate_mi_list.py
import os
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

# 표준 출력 UTF-8 설정 (콘솔 한글 출력 보장)
sys.stdout.reconfigure(encoding='utf-8')

# 작업 경로 및 파일명 설정
workspace = r"c:\Users\han\development\antigraviy\unfi-ing"
file_main = os.path.join(workspace, "2026 도시 조사결과표(한성안).xlsx")
today_str = datetime.now().strftime("%Y-%m-%d")
file_mi_out = os.path.join(workspace, f"2026 도시 조사결과표(한성안)_지역별집계_{today_str}.xlsx")

print("Loading Excel file...")
# 1. 엑셀 파일 데이터 로드
df_main_list = pd.read_excel(file_main, sheet_name="표본점목록")
df_main_res = pd.read_excel(file_main, sheet_name="조사결과표")

# 2. 대상 지역 설정 (수도권 3개 시도: 경기도, 서울특별시, 인천광역시)
target_sidos = ["경기도", "서울특별시", "인천광역시"]

# 제외 대상 시군구 (서울 8개, 경기 5개: 가평군, 남양주시, 양평군, 안성시, 평택시)
exclude_sigungus = {
    "서초구", "강남구", "강동구", "송파구", "성동구", "중랑구", "노원구", "광진구", 
    "가평군", "남양주시", "양평군", "안성시", "평택시"
}

df_list_target = df_main_list[
    df_main_list["시도"].isin(target_sidos) & 
    ~df_main_list["시군구"].isin(exclude_sigungus)
].copy()

# 조사결과표의 시도 컬럼명 확인 ('광역시도' 또는 '시도')
res_sido_col = "광역시도" if "광역시도" in df_main_res.columns else "시도"
df_res_target = df_main_res[
    df_main_res[res_sido_col].isin(target_sidos) & 
    ~df_main_res["시군구"].isin(exclude_sigungus)
].copy()

# 3. 유형 그룹 분류 함수: '가로수' vs '기본' (경관숲, 마을숲, 도시숲, 학교숲)
def get_type_group(val):
    val_str = str(val).strip()
    if val_str == "가로수":
        return "가로수"
    else:
        return "기본"

# 4. 표본점목록 탭: 조사번호에 "필수"가 포함된 표본점 추출 및 유형그룹 부여
df_pilsu = df_list_target[df_list_target["조사번호"].astype(str).str.contains("필수")].copy()
df_pilsu["유형그룹"] = df_pilsu["유형"].apply(get_type_group)

# 5. 조사결과표 탭: 조사가능불가 열의 값이 "조사가능"인 표본점 추출 및 유형그룹 부여
df_res_pos = df_res_target[df_res_target["조사가능불가"].astype(str).str.strip() == "조사가능"].copy()
df_res_pos["유형그룹"] = df_res_pos["유형"].apply(get_type_group)

# 6. 각 시도별 시군구 목록 취합 (가나다순 정렬)
all_sigungu_by_sido = {}
for sido in target_sidos:
    sigungus = set(df_pilsu[df_pilsu["시도"] == sido]["시군구"].dropna().unique()) | \
               set(df_res_pos[df_res_pos[res_sido_col] == sido]["시군구"].dropna().unique())
    # 서울특별시 성동구: 제외 목록 대상 중 가로수 1개 미조사 건으로 집계에 포함
    if sido == "서울특별시":
        sigungus.add("성동구")
    all_sigungu_by_sido[sido] = sorted(list(sigungus))

# 7. 지역별 집계 테이블 데이터 생성
# (1) 전체 총합계 및 시도별 요약 데이터 계산
grand_pilsu_total = 0
grand_res_total = 0
grand_mi_base_total = 0
grand_mi_street_total = 0
grand_mi_total = 0

sido_summary_rows = []
all_district_rows = []

for sido in target_sidos:
    sigungu_list = all_sigungu_by_sido.get(sido, [])
    sido_pilsu_df = df_pilsu[df_pilsu["시도"] == sido]
    sido_res_df = df_res_pos[df_res_pos[res_sido_col] == sido]
    
    sido_pilsu_total = 0
    sido_res_total = 0
    sido_mi_base_total = 0
    sido_mi_street_total = 0
    sido_mi_total = 0
    
    district_rows = []
    for sigungu in sigungu_list:
        if sido == "서울특별시" and sigungu == "성동구":
            # 서울특별시 성동구: 가로수 1개 미조사 수동 추가 반영
            p_cnt = 1
            r_cnt = 0
            rate = 0.0
            mi_base = 0
            mi_street = 1
            mi_cnt = 1
        else:
            p_sub = sido_pilsu_df[sido_pilsu_df["시군구"] == sigungu]
            r_sub = sido_res_df[sido_res_df["시군구"] == sigungu]
            
            # 전체 개수
            p_cnt = len(p_sub)
            r_cnt = len(r_sub)
            rate = (r_cnt / p_cnt) if p_cnt > 0 else 0.0
            
            # 기본 (경관숲, 마을숲, 도시숲, 학교숲) 개수
            p_base = len(p_sub[p_sub["유형그룹"] == "기본"])
            r_base = len(r_sub[r_sub["유형그룹"] == "기본"])
            mi_base = max(0, p_base - r_base)
            
            # 가로수 개수
            p_street = len(p_sub[p_sub["유형그룹"] == "가로수"])
            r_street = len(r_sub[r_sub["유형그룹"] == "가로수"])
            mi_street = max(0, p_street - r_street)
            
            # 총 미조사 필수 개수
            mi_cnt = mi_base + mi_street
        
        sido_pilsu_total += p_cnt
        sido_res_total += r_cnt
        sido_mi_base_total += mi_base
        sido_mi_street_total += mi_street
        sido_mi_total += mi_cnt
        
        district_rows.append({
            "시도": sido,
            "시군구": sigungu,
            "필수 개수": p_cnt,
            "조사완료": r_cnt,
            "진행율": rate,
            "기본 미조사 개수": mi_base,
            "가로수 미조사 개수": mi_street,
            "총 미조사 개수": mi_cnt,
            "row_type": "district"
        })
        
    sido_rate = (sido_res_total / sido_pilsu_total) if sido_pilsu_total > 0 else 0.0
    
    # 시도별 소계 행
    sido_summary_rows.append({
        "시도": sido,
        "시군구": "합계",
        "필수 개수": sido_pilsu_total,
        "조사완료": sido_res_total,
        "진행율": sido_rate,
        "기본 미조사 개수": sido_mi_base_total,
        "가로수 미조사 개수": sido_mi_street_total,
        "총 미조사 개수": sido_mi_total,
        "row_type": "sido_total"
    })
    
    # 누적 전체 총합계
    grand_pilsu_total += sido_pilsu_total
    grand_res_total += sido_res_total
    grand_mi_base_total += sido_mi_base_total
    grand_mi_street_total += sido_mi_street_total
    grand_mi_total += sido_mi_total
    
    all_district_rows.extend(district_rows)

# 전체 총합계 행
grand_rate = (grand_res_total / grand_pilsu_total) if grand_pilsu_total > 0 else 0.0
grand_total_row = {
    "시도": "전체",
    "시군구": "합계",
    "필수 개수": grand_pilsu_total,
    "조사완료": grand_res_total,
    "진행율": grand_rate,
    "기본 미조사 개수": grand_mi_base_total,
    "가로수 미조사 개수": grand_mi_street_total,
    "총 미조사 개수": grand_mi_total,
    "row_type": "grand_total"
}

# 최종 행 순서:
# 1. 타이틀 바로 아래에 '전체 합계' 행 배치
# 2. 이어서 각 3개 시도(경기도, 서울특별시, 인천광역시) 합계 요약 행 배치
# 3. 그 아래에 각 시군구별 상세 행 배치
agg_rows = [grand_total_row] + sido_summary_rows + all_district_rows

print("지역별 집계 생성 완료. 엑셀 파일 스타일링 및 작성 중...")

# 8. openpyxl을 이용한 엑셀 생성 및 스타일링
wb = openpyxl.Workbook()
ws_agg = wb.active
ws_agg.title = "지역별_집계"

# 헤더 작성 (8개 컬럼: 진행율 추가)
headers = [
    "시도", "시군구", "필수 개수", "조사완료", "진행율",
    "기본 미조사 개수", "가로수 미조사 개수", "총 미조사 개수"
]
ws_agg.append(headers)

# 헤더 스타일
header_fill = PatternFill(start_color="FF2F5597", end_color="FF2F5597", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFFFF")

for col in range(1, 9):
    cell = ws_agg.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 배경색 테마
fill_grand = PatternFill(start_color="FFD9E1F2", end_color="FFD9E1F2", fill_type="solid")    # 전체 합계 (소프트 블루그레이)
fill_gyeonggi = PatternFill(start_color="FFFCE4D6", end_color="FFFCE4D6", fill_type="solid") # 경기도 합계 (피치)
fill_seoul = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")    # 서울 합계 (스카이블루)
fill_incheon = PatternFill(start_color="FFE8DDF7", end_color="FFE8DDF7", fill_type="solid")  # 인천 합계 (라벤더)

bold_font = Font(name="맑은 고딕", size=11, bold=True)
normal_font = Font(name="맑은 고딕", size=11, bold=False)
red_bold_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFF0000")

thin_border = Border(
    left=Side(style='thin', color='FFD9D9D9'),
    right=Side(style='thin', color='FFD9D9D9'),
    top=Side(style='thin', color='FFD9D9D9'),
    bottom=Side(style='thin', color='FFD9D9D9')
)

# 데이터 행 입력 및 서식 적용
for i, row_dict in enumerate(agg_rows, start=2):
    ws_agg.cell(row=i, column=1, value=row_dict["시도"])
    ws_agg.cell(row=i, column=2, value=row_dict["시군구"])
    ws_agg.cell(row=i, column=3, value=row_dict["필수 개수"])
    ws_agg.cell(row=i, column=4, value=row_dict["조사완료"])
    ws_agg.cell(row=i, column=5, value=row_dict["진행율"])
    ws_agg.cell(row=i, column=6, value=row_dict["기본 미조사 개수"])
    ws_agg.cell(row=i, column=7, value=row_dict["가로수 미조사 개수"])
    ws_agg.cell(row=i, column=8, value=row_dict["총 미조사 개수"])
    
    row_type = row_dict["row_type"]
    sido_name = row_dict["시도"]
    total_mi = row_dict["총 미조사 개수"]
    
    # 서식 분기: 전체 합계 / 시도별 소계 / 시군구
    if row_type == "grand_total":
        for col in range(1, 9):
            cell = ws_agg.cell(row=i, column=col)
            cell.font = bold_font
            cell.fill = fill_grand
            cell.border = thin_border
            if col in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '0.0%'
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
                
    elif row_type == "sido_total":
        if sido_name == "경기도":
            current_fill = fill_gyeonggi
        elif sido_name == "서울특별시":
            current_fill = fill_seoul
        elif sido_name == "인천광역시":
            current_fill = fill_incheon
        else:
            current_fill = fill_grand
            
        for col in range(1, 9):
            cell = ws_agg.cell(row=i, column=col)
            cell.font = bold_font
            cell.fill = current_fill
            cell.border = thin_border
            if col in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '0.0%'
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
                
    else:
        # 시군구 행: 미조사 필수 개수가 0이면 굵게 + 빨강색 글씨체 적용
        row_font = red_bold_font if total_mi == 0 else normal_font
        for col in range(1, 9):
            cell = ws_agg.cell(row=i, column=col)
            cell.font = row_font
            cell.border = thin_border
            if col in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '0.0%'
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'

# 열 너비 설정 (8개 컬럼)
col_widths = {
    'A': 13.0,  # 시도
    'B': 15.0,  # 시군구
    'C': 13.0,  # 필수 개수
    'D': 14.0,  # 조사완료
    'E': 12.0,  # 진행율
    'F': 18.0,  # 기본 미조사 개수
    'G': 18.0,  # 가로수 미조사 개수
    'H': 16.0   # 총 미조사 개수
}
for col, w in col_widths.items():
    ws_agg.column_dimensions[col].width = w

# 행 높이 설정
ws_agg.row_dimensions[1].height = 28
for r in range(2, len(agg_rows) + 2):
    ws_agg.row_dimensions[r].height = 22

# 파일 저장
wb.save(file_mi_out)
print(f"Excel file successfully generated and saved to: {file_mi_out}")
